import os
import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document


def base_output_path(input_path: str) -> str:
    return os.path.splitext(input_path)[0]


def export_csv(df: pd.DataFrame, base: str) -> str:
    out = base + "_output.csv"
    df.to_csv(out, index=False)
    return out


def export_excel(df: pd.DataFrame, base: str) -> str:
    out = base + "_output.xlsx"
    df.to_excel(out, index=False)
    return out


def export_pdf(df: pd.DataFrame, base: str, last_col: str) -> str:
    out = base + "_output.pdf"
    c = canvas.Canvas(out, pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(200, 800, "Student Report")
    c.setFont("Helvetica", 10)
    y = 770
    for _, row in df.iterrows():
        name = row["Name"] if "Name" in df.columns else "Student"
        line = f"{name} -> {row[last_col]:.2f}"
        c.drawString(100, y, line)
        y -= 15
        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = 800
    c.save()
    return out


def export_docx(df: pd.DataFrame, base: str, last_col: str) -> str:
    out = base + "_output.docx"
    doc = Document()
    doc.add_heading("Student Report", level=1)
    for _, row in df.iterrows():
        name = row["Name"] if "Name" in df.columns else "Student"
        doc.add_paragraph(f"{name} -> {row[last_col]:.2f}")
    doc.save(out)
    return out
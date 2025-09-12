import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document

# Step 1: Import Excel
df = pd.read_excel("students.xlsx")

# Step 2: Detect numeric columns (exclude Total/Average)
numeric_cols = [col for col in df.columns if df[col].dtype in ['int64','float64']]
numeric_cols = [col for col in numeric_cols if col.lower() not in ['total', 'average']]

print("Numeric columns detected:", numeric_cols)

# Step 3: Ask user choice
choice = input("What do you want to calculate? (1 = Addition, 2 = Average): ")

match choice:
    case "1":
        df["Sum"] = df[numeric_cols].sum(axis=1)
        print("✅ Sum calculated for each student")
    case "2":
        df["Average"] = df[numeric_cols].mean(axis=1)
        print("✅ Average calculated for each student")
    case _:
        print("❌ Invalid choice, exiting...")
        exit()

# Step 4: Ask output format
output_choice = input("Export format? (1 = CSV, 2 = Excel, 3 = PDF, 4 = Word): ")

match output_choice:
    case "1":
        df.to_csv("students_output.csv", index=False)
        print("📂 Saved as students_output.csv")
    case "2":
        df.to_excel("students_output.xlsx", index=False)
        print("📂 Saved as students_output.xlsx")
    case "3":
        pdf_file = "students_output.pdf"
        c = canvas.Canvas(pdf_file, pagesize=A4)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(200, 800, "Student Report")

        c.setFont("Helvetica", 10)
        y = 770
        for index, row in df.iterrows():
            line = f"{row['Name']} -> {row[df.columns[-1]]:.2f}"
            c.drawString(100, y, line)
            y -= 15
            if y < 50:  # New page if space ends
                c.showPage()
                c.setFont("Helvetica", 10)
                y = 800
        c.save()
        print("📂 Saved as students_output.pdf")
    case "4":
        doc = Document()
        doc.add_heading("Student Report", level=1)
        for index, row in df.iterrows():
            doc.add_paragraph(f"{row['Name']} -> {row[df.columns[-1]]:.2f}")
        doc.save("students_output.docx")
        print("📂 Saved as students_output.docx")
    case _:
        print("❌ Invalid format selected.")